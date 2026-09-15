"""批量建号的 xlsx 解析与校验。

只负责"文件里的数据本身对不对"，不碰 Keycloak——真正建号复用
KeycloakService.create_user()，跟单人建号走同一条业务逻辑（密码生成规则/
加组/失败回滚），不重新实现一遍。这是现有命令行批量建号脚本（
create_users_from_xlsx.py）的 Web 化，表头格式保持一致，管理员不用重新学一套。
"""

import io
from dataclasses import dataclass
from typing import List, Optional

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.keycloak_service import CREATE_ALLOWED_ROLES

REQUIRED_COLUMNS = ("姓名", "邮箱", "firstName", "lastName", "role")
# 模板/预览用：跟单人建号一样的表头顺序，username 列可选但模板里给出，方便
# 管理员照着填；不用 REQUIRED_COLUMNS 直接拼是因为那个不含 username。
TEMPLATE_HEADER = ("姓名", "邮箱", "firstName", "lastName", "username", "role")
DEFAULT_MAX_ROWS = 1000


@dataclass
class BatchRow:
    row_num: int
    name: str
    email: str
    first_name: str
    last_name: str
    username: str
    role: str
    error: Optional[str] = None

    @property
    def valid(self) -> bool:
        return self.error is None


def parse_xlsx(content: bytes, max_rows: int = DEFAULT_MAX_ROWS) -> List[BatchRow]:
    """解析上传的 xlsx，逐行校验，不合格的行标记 error 但不中断整体解析——
    让管理员在预览页一次性看到所有问题行，不用一行一行改一行一行传。
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"无法解析 xlsx 文件：{exc}") from exc

    ws = wb.active
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise ValueError("xlsx 是空的，没有表头行")
    header = [str(c).strip() if c is not None else "" for c in header_row]
    idx = {name: i for i, name in enumerate(header)}

    missing = [c for c in REQUIRED_COLUMNS if c not in idx]
    if missing:
        raise ValueError(
            f"xlsx 缺少必需的列：{', '.join(missing)}"
            "（表头必须包含 姓名/邮箱/firstName/lastName/role，username 列可选）"
        )
    has_username_col = "username" in idx

    def cell(raw, col):
        i = idx.get(col)
        if i is None or i >= len(raw):
            return ""
        v = raw[i]
        return str(v).strip() if v is not None else ""

    rows: List[BatchRow] = []
    seen_usernames: dict = {}  # username -> 首次出现的 row_num
    seen_emails: dict = {}
    for row_num, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if raw is None or all(c is None for c in raw):
            continue  # 跳过空白行
        if len(rows) >= max_rows:
            raise ValueError(f"xlsx 数据行太多，最多允许 {max_rows} 行")

        name = cell(raw, "姓名")
        email = cell(raw, "邮箱").lower()
        first_name = cell(raw, "firstName")
        last_name = cell(raw, "lastName")
        role = cell(raw, "role")
        username = (cell(raw, "username") if has_username_col else "") or (
            email.split("@", 1)[0] if email else ""
        )
        username = username.lower()

        error = None
        if not email or "@" not in email:
            error = "邮箱为空或格式不对"
        elif not first_name or not last_name:
            error = "firstName/lastName 不能为空"
        elif role not in CREATE_ALLOWED_ROLES:
            error = f"角色 {role!r} 不在支持范围内（批量导入等同于新建账号，只能是 {'/'.join(CREATE_ALLOWED_ROLES)}）"
        elif not username:
            error = "无法确定 username（没填 username 列，邮箱也解析不出来）"
        elif username in seen_usernames:
            error = f"username 跟第 {seen_usernames[username]} 行重复"
        elif email in seen_emails:
            error = f"邮箱跟第 {seen_emails[email]} 行重复"

        if error is None:
            seen_usernames[username] = row_num
            seen_emails[email] = row_num

        rows.append(
            BatchRow(
                row_num=row_num,
                name=name,
                email=email,
                first_name=first_name,
                last_name=last_name,
                username=username,
                role=role,
                error=error,
            )
        )
    return rows


def build_template_xlsx() -> bytes:
    """给管理员下载的批量建号模板：表头 + 一行示例 + role 列加下拉数据验证
    （只能选 CREATE_ALLOWED_ROLES 里那两档），减少手填角色名打错字/打出
    不支持档位的情况——这类错误现在要等上传后在预览页才会被 parse_xlsx
    标红，模板里的下拉能在填表这一步就先挡一部分。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "批量建号"
    ws.append(TEMPLATE_HEADER)
    ws.append(["薛惠坤", "huikun@example.com", "Huikun", "Xue", "", CREATE_ALLOWED_ROLES[0]])

    role_col_letter = "F"  # TEMPLATE_HEADER 第 6 列（role），跟着表头顺序走
    dv = DataValidation(
        type="list",
        formula1='"{}"'.format(",".join(CREATE_ALLOWED_ROLES)),
        allow_blank=False,
        showDropDown=False,  # openpyxl 这个参数名字和实际行为是反的：False 才会显示下拉箭头
        showErrorMessage=True,
    )
    dv.error = f"role 只能是 {'/'.join(CREATE_ALLOWED_ROLES)} 其中一个"
    dv.errorTitle = "角色档位不支持"
    ws.add_data_validation(dv)
    dv.add(f"{role_col_letter}2:{role_col_letter}1000")

    for col_letter, width in zip("ABCDEF", (12, 28, 14, 12, 14, 14)):
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
