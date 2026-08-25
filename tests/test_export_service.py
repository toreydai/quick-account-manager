"""export_service.build_user_list_xlsx()：一键导出用户清单。只测生成的
xlsx 内容本身，不碰 Keycloak/QuickSight——那两部分的数据来源
（_fetch_users/_build_subscription_status_map）已经在别处测过或者是纯
HTTP 调用薄封装，这里只关心"给一份 QuickUser 列表 + 订阅状态 map，
生成的表对不对"。
"""

import io

import openpyxl

from app.services.export_service import HEADER, build_user_list_xlsx
from app.services.keycloak_service import QuickUser


def make_user(**overrides):
    defaults = dict(
        id="u1", username="huikun", email="huikun@example.com",
        first_name="Huikun", last_name="Xue", enabled=True, role="作者版",
    )
    defaults.update(overrides)
    return QuickUser(**defaults)


def test_header_matches_expected_columns():
    data = build_user_list_xlsx([], {})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == list(HEADER)


def test_row_contains_expected_fields():
    users = [make_user()]
    data = build_user_list_xlsx(users, {"u1": "已生效"})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert row == ["huikun", "huikun@example.com", "Huikun", "Xue", "作者版", "启用中", "已生效"]


def test_disabled_user_shown_as_disabled():
    users = [make_user(enabled=False)]
    data = build_user_list_xlsx(users, {})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert row[5] == "已停用"


def test_user_without_role_shows_placeholder_and_no_subscription_status():
    """不在任何 quick-* 组里的人，角色档位显示占位文案，订阅状态直接是
    "—"——不该去 status_map 里查一个本来就不存在订阅这回事的人。"""
    users = [make_user(role=None)]
    data = build_user_list_xlsx(users, {"u1": "已生效"})  # 就算 map 里意外有值也不该被用上
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert row[4] == "（未分配档位）"
    assert row[6] == "—"


def test_user_with_role_but_missing_from_status_map_shows_dash():
    """有角色但 status_map 里没有这个 user_id（比如查询失败/超时被跳过），
    应该显示"—"，不能让 KeyError 把整个导出搞挂。"""
    users = [make_user()]
    data = build_user_list_xlsx(users, {})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert row[6] == "—"


def test_multiple_users_preserve_order():
    users = [make_user(id="u1", username="a"), make_user(id="u2", username="b")]
    data = build_user_list_xlsx(users, {})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    usernames = [row[0].value for row in ws.iter_rows(min_row=2, max_row=3)]
    assert usernames == ["a", "b"]
