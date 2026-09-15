"""batch_import.parse_xlsx() 的解析/校验规则。只测解析本身，不碰 Keycloak——
真正建号复用 KeycloakService.create_user()，那部分逻辑已经在
test_create_user_flow.py 里覆盖了。
"""

import io

import openpyxl
import pytest

from app.services.batch_import import build_template_xlsx, parse_xlsx
from app.services.keycloak_service import CREATE_ALLOWED_ROLES


def make_xlsx(header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADER = ["姓名", "邮箱", "firstName", "lastName", "role"]
HEADER_WITH_USERNAME = ["姓名", "邮箱", "firstName", "lastName", "username", "role"]


def test_parses_valid_rows_and_derives_username_from_email():
    content = make_xlsx(HEADER, [["薛惠坤", "huikun@example.com", "Huikun", "Xue", "作者版"]])
    rows = parse_xlsx(content)

    assert len(rows) == 1
    assert rows[0].valid
    assert rows[0].username == "huikun"
    assert rows[0].role == "作者版"


def test_explicit_username_column_overrides_email_derived_default():
    content = make_xlsx(
        HEADER_WITH_USERNAME,
        [["薛惠坤", "huikun@example.com", "Huikun", "Xue", "hkxue", "作者版"]],
    )
    rows = parse_xlsx(content)

    assert rows[0].username == "hkxue"


def test_missing_required_column_raises():
    content = make_xlsx(["姓名", "邮箱"], [["薛惠坤", "huikun@example.com"]])
    with pytest.raises(ValueError, match="缺少必需的列"):
        parse_xlsx(content)


def test_invalid_role_marks_row_error_but_keeps_parsing_other_rows():
    content = make_xlsx(
        HEADER,
        [
            ["A", "a@example.com", "A", "A", "不存在的档位"],
            ["B", "b@example.com", "B", "B", "作者版"],
        ],
    )
    rows = parse_xlsx(content)

    assert len(rows) == 2
    assert not rows[0].valid
    assert "不在支持范围内" in rows[0].error
    assert rows[1].valid


def test_duplicate_username_within_file_marked_invalid():
    content = make_xlsx(
        HEADER,
        [
            ["A", "same@example.com", "A", "A", "作者版"],
            ["B", "same@example.com", "B", "B", "作者版"],
        ],
    )
    rows = parse_xlsx(content)

    assert rows[0].valid
    assert not rows[1].valid
    assert "重复" in rows[1].error


def test_blank_rows_are_skipped():
    content = make_xlsx(HEADER, [[None, None, None, None, None], ["A", "a@example.com", "A", "A", "作者版"]])
    rows = parse_xlsx(content)

    assert len(rows) == 1
    assert rows[0].email == "a@example.com"


def test_missing_email_marks_row_invalid():
    content = make_xlsx(HEADER, [["A", "", "A", "A", "作者版"]])
    rows = parse_xlsx(content)

    assert not rows[0].valid
    assert "邮箱" in rows[0].error


# ---------- 2026-08-24：批量导入等同于新建账号，角色范围收窄回两档 ----------


def test_role_outside_create_allowed_roles_marks_row_invalid():
    """"阅读版"是 ROLE_TO_GROUP 里真实存在的档位（改档位能用），但不在
    CREATE_ALLOWED_ROLES 里——批量导入是"新建账号"的一种，跟单人建号走
    同一套角色范围，这种行要被标红，不是"不存在的档位"那种完全无效值。"""
    content = make_xlsx(HEADER, [["A", "a@example.com", "A", "A", "阅读版"]])
    rows = parse_xlsx(content)

    assert not rows[0].valid
    assert "不在支持范围内" in rows[0].error


@pytest.mark.parametrize("role", CREATE_ALLOWED_ROLES)
def test_each_create_allowed_role_is_accepted(role):
    content = make_xlsx(HEADER, [["A", "a@example.com", "A", "A", role]])
    rows = parse_xlsx(content)

    assert rows[0].valid


# ---------- 模板下载 ----------


def test_build_template_xlsx_has_expected_header_and_example_row():
    data = build_template_xlsx()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == ["姓名", "邮箱", "firstName", "lastName", "username", "role"]

    example_role = ws.cell(row=2, column=6).value
    assert example_role in CREATE_ALLOWED_ROLES


def test_template_role_column_has_dropdown_limited_to_allowed_roles():
    data = build_template_xlsx()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active

    validations = ws.data_validations.dataValidation
    assert len(validations) == 1
    dv = validations[0]
    assert dv.type == "list"
    for role in CREATE_ALLOWED_ROLES:
        assert role in dv.formula1


def test_template_example_row_parses_as_valid():
    """模板生成的示例行本身也要能通过 parse_xlsx 校验，不能自己都不合规。"""
    data = build_template_xlsx()
    rows = parse_xlsx(data)

    assert len(rows) == 1
    assert rows[0].valid


def test_parse_xlsx_rejects_too_many_rows():
    content = make_xlsx(
        HEADER,
        [
            ["A", "a@example.com", "A", "A", "作者版"],
            ["B", "b@example.com", "B", "B", "作者版"],
        ],
    )

    with pytest.raises(ValueError, match="最多允许 1 行"):
        parse_xlsx(content, max_rows=1)
